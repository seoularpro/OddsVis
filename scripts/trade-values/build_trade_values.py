#!/usr/bin/env python3
"""Build the nine Scoring x League-size trade value workbooks from the Half PPR 10-team baseline.

    python3 scripts/trade-values/build_trade_values.py                 # baseline = live published sheet
    python3 scripts/trade-values/build_trade_values.py --baseline my.xlsx   # baseline = a local workbook/CSV in the sheet layout
    python3 scripts/trade-values/build_trade_values.py --refresh-receptions --week 3   # re-read reception lines from BettingProsFiles

Inputs (all next to this script):
    params.json       every constant of the adjustment math
    receptions.json   BettingPros receptions line per player (snapshot; refresh with --refresh-receptions)
Output (all under TradeValueSheets/):
    TradeValues_<Scoring>_<N>team.xlsx (nine files) and TradeValues_AllCombinations.xlsx
    trade-values.json            every combination in one file, plus a byPlayer index (for the Chrome extension)
    json/TradeValues_<Scoring>_<N>team.json   one file per combination (what the site's dropdowns load)
The GitHub Action in .github/workflows/tradeValues.yml runs this on a schedule and commits the results.

The math, in order:
  1. Scoring (zero-sum).  Players whose receptions line is at least `significant_reception_margin` above their
     position's replacement level move by  k * delta * (line - replacement),  k = value_points_per_fantasy_point,
     delta = the format's reception_delta (Full PPR +0.5, Standard -0.5, Half PPR 0).  The total of those moves is
     taken back in equal amounts from every other player, so the pool's total value is unchanged.  Nobody goes
     below `minimum_value`; what a player cannot absorb is spread over the remaining absorbers.
     Replacement level = mean line of baseline players valued <= `replacement_value_cutoff` at the position.
     Players with no line use the mean line of position peers within `peer_value_window` value points.
  2. League size.  final = top * (scored / baseline_top) ^ exponent, rounded, floored at minimum_value.
Requires: openpyxl (python3 -m pip install --user openpyxl).
"""
import argparse, csv, datetime, glob, html, json, os, re, sys, urllib.request
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
POS_ORDER = {'RB': 0, 'WR': 1, 'TE': 2, 'QB': 3}


def norm(n):
    return re.sub(r'\.', '', n or '').replace(' Jr', '').replace(' jr', '').replace(' Sr', '').replace(' sr', '').strip()


# ----------------------------------------------------------------------------- baseline
def grid_from_published_sheet(url):
    sheet_url = url.replace('/pubhtml?', '/pubhtml/sheet?headers=false&')
    with urllib.request.urlopen(sheet_url, timeout=30) as r:
        h = r.read().decode('utf-8')
    rows = []
    for tr in re.findall(r'<tr[^>]*>(.*?)</tr>', h, re.S):
        tds = re.findall(r'<td[^>]*>(.*?)</td>', tr, re.S)
        if tds:
            rows.append([html.unescape(re.sub(r'<[^>]+>', '', t)).strip() for t in tds])
    return rows


def grid_from_file(path):
    if path.lower().endswith('.csv'):
        with open(path, newline='') as f:
            return [[(c or '').strip() for c in row] for row in csv.reader(f)]
    wb = load_workbook(path, data_only=True)
    ws = wb['Trade Values'] if 'Trade Values' in wb.sheetnames else wb[wb.sheetnames[0]]
    return [['' if c.value is None else str(c.value).strip() for c in row] for row in ws.iter_rows()]


def parse_grid(grid):
    """The sheet layout: a position label row, then 'Trade Value | Player Name | Trend | ...', then data."""
    last = max((i for r in grid for i, c in enumerate(r) if c), default=-1)
    rows = [r[:last + 1] + [''] * (last + 1 - len(r)) for r in grid]
    while rows and not any(rows[-1]):
        rows.pop()
    hdr = next(i for i, r in enumerate(rows) if any(c.lower() == 'trade value' for c in r))
    posrow, header = rows[hdr - 1], rows[hdr]
    cols = []
    for c, cell in enumerate(header):
        if cell.lower() == 'trade value':
            cols.append({'kind': 'value'})
        elif cell.lower() == 'player name':
            lab = next((posrow[k].upper() for k in range(c, -1, -1) if posrow[k].upper() in POS_ORDER), None)
            cols.append({'kind': 'name', 'pos': lab})
        elif cell.lower() == 'trend':
            cols.append({'kind': 'trend'})
        else:
            cols.append({'kind': 'other'})
    vcol = next(i for i, c in enumerate(cols) if c['kind'] == 'value')
    players, order = [], 0
    for r in rows[hdr + 1:]:
        try:
            v = float(r[vcol])
        except ValueError:
            continue
        for i, c in enumerate(cols):
            if c['kind'] == 'name' and r[i]:
                players.append({'name': r[i], 'pos': c['pos'], 'value': v, 'order': order}); order += 1
    return {'posrow': posrow, 'header': header, 'columns': cols, 'players': players}


# ----------------------------------------------------------------------------- receptions
def refresh_receptions(year, week):
    """Consensus receptions line (market 104) from the newest, carry, and first BettingPros files of the week."""
    d = os.path.join(REPO, 'BettingProsFiles'); prefix = '' if year == 2023 else str(year)
    try:
        last = int(open(os.path.join(d, f'{prefix}lastIndex{week}.txt')).read().strip())
    except OSError:
        last = max((int(m.group(1)) for f in glob.glob(os.path.join(d, f'{prefix}week{week}*'))
                    for m in [re.fullmatch(rf'{prefix}week{week}(\d+)', os.path.basename(f))] if m), default=0)
    sources = [(os.path.join(d, f'{prefix}week{week}{last}'), f'week{week} latest'),
               (os.path.join(d, f'{prefix}carry{week}.json'), f'week{week} carry'),
               (os.path.join(d, f'{prefix}week{week}0'), f'week{week} first')]
    out = {}
    for path, src in sources:
        try:
            data = json.load(open(path))
        except (OSError, ValueError):
            continue
        for p in data.get('props', []):
            if p.get('market_id') != 104:
                continue
            name = p['participant']['name']; line = (p.get('over') or {}).get('consensus_line') or (p.get('over') or {}).get('line')
            if line is not None and norm(name) not in out:
                out[norm(name)] = {'name': name, 'pos': p['participant'].get('player', {}).get('position'), 'line': line, 'source': src}
    return {'_note': f'BettingPros consensus receptions line per player (market 104) from the {year} week {week} files.',
            'season': year, 'week': week, 'players': dict(sorted(out.items()))}


# ----------------------------------------------------------------------------- the math
def attach_receptions(players, recs, P):
    for p in players:
        r = recs.get(norm(p['name']))
        p['rec'], p['est'] = (r['line'], False) if r else (None, True)
    repl = {}
    for pos in POS_ORDER:
        xs = [p['rec'] for p in players if p['pos'] == pos and p['value'] <= P['replacement_value_cutoff'] and p['rec'] is not None]
        repl[pos] = sum(xs) / len(xs) if xs else 0.0
    for p in players:
        if p['rec'] is None:
            peers = [q['rec'] for q in players if q['pos'] == p['pos'] and q['rec'] is not None and abs(q['value'] - p['value']) <= P['peer_value_window']] \
                or [q['rec'] for q in players if q['pos'] == p['pos'] and q['rec'] is not None]
            p['rec'] = round(sum(peers) / len(peers), 1) if peers else 0.0
    return repl


def scored_values(players, repl, delta, P):
    k, floor = P['value_points_per_fantasy_point'], P['minimum_value']
    scored = {p['name']: float(p['value']) for p in players}
    if delta == 0:
        return scored
    sig = lambda p: p['rec'] >= repl[p['pos']] + P['significant_reception_margin']
    total = 0.0
    for p in players:
        if sig(p):
            adj = k * delta * (p['rec'] - repl[p['pos']]); scored[p['name']] += adj; total += adj
    remaining, pool = -total, [p for p in players if not sig(p)]
    while abs(remaining) > 1e-9 and pool:
        share, nxt = remaining / len(pool), []
        for p in pool:
            room = scored[p['name']] - floor
            take = share if share >= 0 or -share <= room else -room
            scored[p['name']] += take; remaining -= take
            if not (share < 0 and -share > room):
                nxt.append(p)
        pool = nxt
    return scored


def league_value(scored, top, exponent, baseline_top, floor):
    return max(floor, round(top * (max(float(floor), scored) / baseline_top) ** exponent))


# ----------------------------------------------------------------------------- workbooks
F = lambda **k: Font(name='Arial', **k)
THIN = Side(style='thin', color='D9D9D9'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
POS_FILL = {'RB': '10B981', 'WR': '0EA5E9', 'TE': 'F59E0B', 'QB': 'F43F5E'}; HEAD_FILL = PatternFill('solid', fgColor='EFEFEF')


def write_grid(ws, base, rows, note):
    cols = base['columns']
    ws.append([c.upper() if c else None for c in base['posrow']])
    ws.append([c or None for c in base['header']])
    ws.append([None] * len(cols))
    for p in sorted(rows, key=lambda p: (-p['final'], POS_ORDER[p['pos']], -p['value'], p['order'])):
        ws.append([p['final'] if c['kind'] == 'value' else (p['name'] if c['kind'] == 'name' and c['pos'] == p['pos'] else None) for c in cols])
    for c in ws[1]:
        if c.value in POS_FILL:
            c.font = F(bold=True, color='FFFFFF', size=10); c.fill = PatternFill('solid', fgColor=POS_FILL[c.value]); c.alignment = Alignment(horizontal='center')
    for c in ws[2]:
        c.font = F(bold=True, size=10); c.fill = HEAD_FILL; c.border = BORDER
    for r in range(3, ws.max_row + 1):
        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i); cell.border = BORDER
            cell.font = F(bold=True, size=11) if col['kind'] == 'value' else F(size=10)
            if col['kind'] == 'value':
                cell.alignment = Alignment(horizontal='right')
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = {'value': 13, 'name': 24, 'trend': 8}.get(col['kind'], 12)
    ws.freeze_panes = 'B3'; ws['A1'].comment = Comment(note, 'VegasLytics')


def write_readme(wb, P, label, size, delta, top, exponent, repl, rows, baseline_desc, rec_desc):
    r = wb.create_sheet('README'); k = P['value_points_per_fantasy_point']
    movers = sorted(rows, key=lambda p: -abs(p['final'] - p['value']))[:8]
    lines = [
        f'Trade Values — {label}, {size}-team league',
        '',
        f'Baseline: {baseline_desc} ({P["baseline_label"]}), {len(rows)} players, top value {max(p["value"] for p in rows):g}.',
        'Derived in two steps (plain numbers; one row per player, sorted by value). Generated by scripts/trade-values/build_trade_values.py.',
        '',
        '1) Scoring (zero-sum). Players with a real receiving role gain value in Full PPR and lose the same amount in Standard,',
        '   relative to Half PPR; the total of those moves is taken back evenly (equal amount per player) from everyone else,',
        f'   so the pool\'s total trade value is unchanged. No player drops below {P["minimum_value"]}.',
        f'   receiver move = {k:g} x delta x (receptions/game - replacement receptions/game at the position), delta = {delta:+.1f} for this file.',
        f'   A receiving role = receptions/game at least {P["significant_reception_margin"]:g} above the position\'s replacement level.',
        f'   receptions/game = {rec_desc}; players with no line use the mean line of position peers within {P["peer_value_window"]} value points.',
        f'   Replacement level = average line of baseline players valued {P["replacement_value_cutoff"]} or less at the position: '
        + ', '.join(f'{k2} {v:.1f}' for k2, v in repl.items()) + '.',
        '',
        '2) League size. Bigger leagues widen the gap between stars and depth; smaller leagues compress it.',
        f'   final = top x (value / baseline top) ^ exponent, with top = {top} and exponent = {exponent} for this file ('
        + ', '.join(f'{s} teams {v["top"]}/{v["exponent"]}' for s, v in P['league_sizes'].items()) + ').',
        '',
        'Largest moves vs. the baseline in this file: ' + ('none' if all(p['final'] == p['value'] for p in rows) else
                                                          ', '.join(f"{p['name']} {p['value']:g}->{p['final']}" for p in movers)),
    ]
    for i, t in enumerate(lines, start=1):
        r.cell(row=i, column=1, value=t).font = F(bold=(i == 1), size=12 if i == 1 else 10)
    r.column_dimensions['A'].width = 140


# ----------------------------------------------------------------------------- JSON (for the site and the Chrome extension)
JSON_SCHEMA_VERSION = 1


def num(v):
    """Whole numbers as ints in the JSON (70, not 70.0)."""
    return int(v) if float(v).is_integer() else v


def sheet_json(base, key, size, sc, rows):
    """One combination: the grid exactly as the sheet lays it out, plus a flat sorted player list."""
    cols = base['columns']
    ordered = sorted(rows, key=lambda p: (-p['final'], POS_ORDER[p['pos']], -p['value'], p['order']))
    grid = [[p['final'] if c['kind'] == 'value' else (p['name'] if c['kind'] == 'name' and c['pos'] == p['pos'] else '') for c in cols]
            for p in ordered]
    return {
        'key': f'{key}_{size}', 'scoring': key, 'scoringLabel': sc['label'], 'leagueSize': int(size),
        'xlsx': f'TradeValues_{key}_{size}team.xlsx',
        'positionRow': [c.upper() if c else '' for c in base['posrow']],
        'header': list(base['header']),
        'rows': grid,
        'players': [{'name': p['name'], 'pos': p['pos'], 'value': p['final'], 'baselineValue': num(p['value'])} for p in ordered],
    }


def write_json(out_dir, base, P, R, baseline_desc, sheets, repl, generated_at):
    """trade-values.json = everything in one file; json/<name>.json = one file per combination."""
    players = base['players']
    by_player = {}
    for p in players:
        by_player[p['name']] = {'pos': p['pos'], 'baselineValue': num(p['value']), 'values': {}}
    for key, sh in sheets.items():
        for p in sh['players']:
            by_player[p['name']]['values'][key] = p['value']
    combined = {
        'schemaVersion': JSON_SCHEMA_VERSION,
        'generatedAt': generated_at,
        'baseline': {'label': P['baseline_label'], 'source': baseline_desc, 'sheetUrl': P['baseline_sheet_url'],
                     'playerCount': len(players), 'topValue': num(max(p['value'] for p in players))},
        'receptions': {'season': R['season'], 'week': R['week'],
                       'replacementPerGame': {k: round(v, 2) for k, v in repl.items()}},
        'params': {k: v for k, v in P.items() if not k.startswith('_')},
        'columns': list(base['header']),
        'sheetKeys': list(sheets.keys()),
        'sheets': sheets,
        'byPlayer': by_player,
    }
    with open(os.path.join(out_dir, 'trade-values.json'), 'w') as f:
        json.dump(combined, f, indent=1)
    jdir = os.path.join(out_dir, 'json'); os.makedirs(jdir, exist_ok=True)
    for key, sh in sheets.items():
        with open(os.path.join(jdir, f'TradeValues_{sh["scoring"]}_{sh["leagueSize"]}team.json'), 'w') as f:
            json.dump({'schemaVersion': JSON_SCHEMA_VERSION, 'generatedAt': generated_at, **sh}, f, indent=1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--baseline', help='local .xlsx/.csv in the sheet layout (default: the live published sheet in params.json)')
    ap.add_argument('--out', default=os.path.join(REPO, 'TradeValueSheets'))
    ap.add_argument('--refresh-receptions', action='store_true', help='re-read reception lines from BettingProsFiles and save receptions.json')
    ap.add_argument('--year', type=int, default=None); ap.add_argument('--week', type=int, default=None)
    a = ap.parse_args()

    P = json.load(open(os.path.join(HERE, 'params.json')))
    rec_path = os.path.join(HERE, 'receptions.json')
    if a.refresh_receptions:
        old = json.load(open(rec_path)) if os.path.exists(rec_path) else {}
        R = refresh_receptions(a.year or old.get('season', 2026), a.week or old.get('week', 1))
        json.dump(R, open(rec_path, 'w'), indent=1); print(f'receptions.json refreshed: {len(R["players"])} lines from {R["season"]} week {R["week"]}')
    R = json.load(open(rec_path))
    rec_desc = f'BettingPros consensus receptions line from the {R["season"]} week {R["week"]} props'

    if a.baseline:
        base, baseline_desc = parse_grid(grid_from_file(a.baseline)), f'local file {os.path.basename(a.baseline)}'
    else:
        base, baseline_desc = parse_grid(grid_from_published_sheet(P['baseline_sheet_url'])), 'the live Trade Values Google Sheet'
    players = base['players']
    if len(players) < P.get('minimum_player_count', 50):
        sys.exit(f'refusing to publish: baseline has only {len(players)} players (expected at least {P.get("minimum_player_count", 50)})')
    repl = attach_receptions(players, R['players'], P)
    baseline_top = max(p['value'] for p in players)
    print(f'baseline: {baseline_desc}, {len(players)} players, top {baseline_top:g}; replacement rec/game', {k: round(v, 2) for k, v in repl.items()},
          '; estimated lines:', sum(p['est'] for p in players))

    os.makedirs(a.out, exist_ok=True)
    generated_at = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace('+00:00', 'Z')
    combined = Workbook(); combined.remove(combined.active); sheets = {}
    for key, sc in P['scoring'].items():
        for size, lg in P['league_sizes'].items():
            scored = scored_values(players, repl, sc['reception_delta'], P)
            rows = [{**p, 'scored': scored[p['name']], 'final': league_value(scored[p['name']], lg['top'], lg['exponent'], baseline_top, P['minimum_value'])} for p in players]
            sheets[f'{key}_{size}'] = sheet_json(base, key, size, sc, rows)
            note = f'{sc["label"]}, {size}-team. Derived from {baseline_desc} ({P["baseline_label"]}); see README.'
            wb = Workbook(); ws = wb.active; ws.title = 'Trade Values'; write_grid(ws, base, rows, note)
            write_readme(wb, P, sc['label'], size, sc['reception_delta'], lg['top'], lg['exponent'], repl, rows, baseline_desc, rec_desc)
            path = os.path.join(a.out, f'TradeValues_{key}_{size}team.xlsx'); wb.save(path); print('wrote', os.path.relpath(path, REPO))
            write_grid(combined.create_sheet(f'{sc["label"]} {size}-team'), base, rows, note)
    rd = combined.create_sheet('README', 0)
    rd['A1'] = 'All nine Scoring x League-size grids, one tab each, one row per player. See any individual workbook README for the method.'
    rd['A1'].font = F(size=10); rd.column_dimensions['A'].width = 140
    combined.save(os.path.join(a.out, 'TradeValues_AllCombinations.xlsx')); print('wrote', os.path.relpath(os.path.join(a.out, 'TradeValues_AllCombinations.xlsx'), REPO))
    write_json(a.out, base, P, R, baseline_desc, sheets, repl, generated_at)
    print('wrote', os.path.relpath(os.path.join(a.out, 'trade-values.json'), REPO), 'and', len(sheets), 'files under', os.path.relpath(os.path.join(a.out, 'json'), REPO))


if __name__ == '__main__':
    main()
